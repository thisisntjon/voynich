"""S64: where does the Naibbe cipher (Greshko 2025) land in the s55 wrapper-free taxonomy, next to Voynich?

QUESTION (frozen): in the Naibbe cipher the table (and with it the conserved prefix family) is a random card draw
per plaintext letter and the plaintext letter sets the suffix, so its prefix stream should be close to orderless.
Voynich's prefix stream is more predictable than any code or generator built so far (CONSTRAINTS-MAP dot 15;
s55 f4_prefix_gain ~0.095-0.100).  Does Naibbe ciphertext land where Voynich does on s55's features?

PRE-REGISTRATION: frozen in s64_naibbe_taxonomy.prereg.sha256 (sha256sum format: this file and every local module it
imports, incl. s55_code_vs_generator.py and its imports) BEFORE any real-text scoring.  Before the freeze the script
ran only with --synthetic (synthetic tables, synthetic plaintext, s55's synthetic fixture corpus; reads no real data;
writes only runs/s64/fixture.txt).

NAIBBE ROUTE [CHOICE, recorded]: the only Naibbe lines on disk from Greshko's own tables are ONE cached encryption in
his workbook (data/naibbe_zenodo_17219445/Naibbe_Cipher.xlsx, sheet 'analysis' col C: 20,923 tokens of his
'Four-Text Sampler', no page/line structure; used by c4).  One text at 21k tokens of a different plaintext cannot form
a class matched to s55 (38k-token texts, s55 plaintext languages, 3 seeds x 3 plaintexts).  So the encipherment is
RE-IMPLEMENTED from Greshko's own tables, read as data (values only, openpyxl data_only) from the same workbook:
  unigram tables  : sheet unigram_lookup, key col E 'unigram_<table>_<letter>' -> col K ('manual')   (= S1 Table 2)
  bigram affixes  : sheet presuf_lookup, key col E '<table>_<letter>_prefix|suffix' -> col F
  deck            : sheet 'deck and spacing', the 78-card row: cumulative thresholds alpha, beta1, beta2, beta3,
                    gamma1, gamma2 (the workbook's default deck, the one behind his cached run)
  spacing         : the 'Standard (47.2%-52.8%-0%)' row: P(unigram chunk) = 0.4722..., else bigram; no trigrams
  procedure       : exactly the workbook's sheet 'encrypter' (rows 15+): the plaintext is a letter stream without
                    spaces; repeatedly draw chunk length (1 if U < p_uni else 2), take the next chunk; draw a table
                    independently per chunk letter (U against the cumulative deck thresholds); unigram chunk ->
                    unigram table entry of (table1, letter); bigram chunk -> prefix(table1, letter1) +
                    suffix(table2, letter2).  One chunk = one ciphertext token.
  plaintext letters: NFKD, lower case, combining marks dropped, j->i, k->c, w->v [CHOICE: the tables have the 23
                    letters a-i, l-v, x-z], every other character dropped (Latin cremma's abbreviation signs are lost).
Plaintexts [CHOICE]: three of s55's CODE plaintext languages, the two Greshko designed for plus one other Romance:
  latin_cremma, italian_dante, occitan_croisade = the word streams of runs/s49/lang_<nm>.meta.json (read-only), joined
  without spaces; cycled if exhausted (flagged).  Seeds 1, 2, 3 per plaintext: random.Random(64000 + 100*i + s).
Size and layout: each text = exactly 38,000 tokens poured in order into the page/line skeleton of
  runs/s49/code_k4_<nm>.meta.json (the ZL3b 38k skeleton, page ids and Voynich section labels, as s55's CODE class).
Reported-only extra (never in a class): naibbe_greshko_cached = Greshko's cached 20,923 tokens poured into the first
  20,923 skeleton slots of code_k4_latin_cremma.

SCORING: s55's pipeline imported verbatim (import s55_code_vs_generator as S55: S55.TextData(...).measure(),
  S55.prep, S55.nc_fit_predict, S55.lr_fit, S55.lr_predict, S55.voynich_pages, S55.load_sidecar).  s55's 12 texts
  (CODE 6, GEN 6, runs/s49/labels.json) and Voynich ZL3b / RF1b / IT2a are re-scored and checked against
  runs/s55/features.csv (max |diff| printed).

PREDICTIONS (fixed):
  P1 (per transliteration T): NAIBBE f4_prefix_gain is below Voynich_T's.  Interval: page subsampling of the held-out
     pages, 80 % of them (round(0.8 n)) WITHOUT replacement, 500 draws, np.random.default_rng(64064); per draw f4 of
     each NAIBBE text from its per-page gain rows (S55.gain_stat), class value = mean of the 9 texts.  P1 holds for T
     iff the 97.5th percentile of the class-mean draws < Voynich_T's point f4 AND every NAIBBE text's own 97.5th
     percentile < Voynich_T's point f4.
  P2 (per T): Voynich_T lies outside the NAIBBE class region.  s55 had only a two-class nearest-centroid / LR geometry
     and no outside-class test, so its geometry is extended [DEVIATION, stated before the run]: z-space of s55
     (S55.prep, mean/sd of the 12 s55 texts, all 8 features); centroids CODE, GEN (s55's 12) and NAIBBE (the 9).
     Voynich_T is outside iff (a) its nearest of the three centroids is not NAIBBE AND (b) its distance to the NAIBBE
     centroid exceeds R = the largest leave-one-out distance of a NAIBBE text to the centroid of the other 8.
  VERDICT: "random-table-per-token cipher excluded" if P1 and P2 hold on all three transliterations; "not excluded"
     if either fails on any.  If the planted round-trip (below) fails, the verdict is "INVALID (implementation does
     not round-trip)" instead.
REPORTED, NOT GATING: every s55 feature for each NAIBBE text, the NAIBBE class mean (sd), s55 class means and Voynich;
  80 %-subsample intervals for f1/f3/f4/f5a/f5b of NAIBBE; Voynich 80 %-subsample draws (500, 20 null draws each, as
  s55's bootstrap but without replacement) and the share of draws meeting P2; s55 LR P(CODE) of each NAIBBE text.
PLANTED CHECK: round-trip on every NAIBBE text: decipher with the tables plus the per-token key record (chunk length
  and card draws) -> must equal the plaintext letters exactly (gating, see VERDICT).  Also reported: decoding with the
  tables plus chunk lengths only, and with the tables only (share of tokens with a unique reading).
FIDELITY (reported): Greshko reports repeat rate 0.18 +- 0.04 % (S1.6) and MATTR 0.56 (S1.9); his workbook defines
  the repeat rate as (# t in tokens 2..10000 with token_t == token_t-1) / 10000.  Computed here with the same repeat
  definition and MATTR = mean type/token ratio over all 1000-token sliding windows [CHOICE: the workbook's sampled
  MATTR cells are not reproducible from cached values] on (i) Greshko's cached output, (ii) this implementation on his
  own Four-Text Sampler plaintext (sheet 'text dropdown'), seeds 1-3, (iii) the 9 NAIBBE texts.
Constraints: python -I, one CPU process, no new data; writes ONLY results_s64.txt and runs/s64/; runs/s49, runs/s55
  read-only.  Run: python -I s64_naibbe_taxonomy.py      Pre-freeze: python -I s64_naibbe_taxonomy.py --synthetic
"""
from __future__ import annotations

import csv
import hashlib
import json
import random
import sys
import time
import unicodedata
from collections import Counter, OrderedDict
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:  # noqa: BLE001
        pass

import s55_code_vs_generator as S55  # noqa: E402

SELF = ROOT / "s64_naibbe_taxonomy.py"
STUB = ROOT / "s64_naibbe_taxonomy.prereg.sha256"
XLSX = ROOT / "data" / "naibbe_zenodo_17219445" / "Naibbe_Cipher.xlsx"
S49 = ROOT / "runs" / "s49"
S55_FEAT = ROOT / "runs" / "s55" / "features.csv"
RUN_DIR = ROOT / "runs" / "s64"
SYNTH = "--synthetic" in sys.argv
OUT = RUN_DIR / "fixture.txt" if SYNTH else ROOT / "results_s64.txt"
FEATS = S55.FEATS
TABLES = ["alpha", "beta1", "beta2", "beta3", "gamma1", "gamma2"]
LETTERS23 = "abcdefghilmnopqrstuvxyz"
PLAINS = ["latin_cremma", "italian_dante", "occitan_croisade"]
SEEDS = [1, 2, 3]
N_TOK = 38000
B_SUB = 500
SUB_FRAC = 0.8
SUB_SEED = 64064
N_SHUF_SUB = 20
WALL = 30 * 60
TRANSLITS = list(S55.TRANSLITS)

log: list[str] = []


def emit(s=""):
    print(s, flush=True)
    log.append(s)


def flush():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text("\n".join(log) + "\n", encoding="utf-8", newline="\n")


def sha256(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def stub_match():
    if not STUB.exists():
        return "NO STUB"
    want = {}
    for ln in STUB.read_text(encoding="utf-8").splitlines():
        ln = ln.strip()
        if not ln or ln.startswith("#") or ln.startswith("FAILED") or ln.startswith("OK"):
            continue
        h, f = ln.split(None, 1)
        want[f.lstrip("*").strip()] = h
    bad = [f for f, h in want.items() if not (ROOT / f).exists() or sha256(ROOT / f) != h]
    missing = [m for m in S55.local_modules() + [SELF.name] if m not in want]
    ok = bool(want) and not bad and not missing
    return "True" if ok else f"False (mismatch: {', '.join(bad) or '-'}; loaded but not in stub: {', '.join(missing) or '-'})"


# ------------------------------------------------------------------ Naibbe tables and encipherment
def load_greshko():
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    uni, pre, suf = {}, {}, {}
    for r in wb["unigram_lookup"].iter_rows(min_row=2, max_row=139, max_col=11, values_only=True):
        _, t, a = str(r[4]).split("_")  # 'unigram_<table>_<letter>'
        uni[(t, a)] = str(r[10])
    for r in wb["presuf_lookup"].iter_rows(min_row=2, max_row=277, max_col=6, values_only=True):
        t, a, kind = str(r[4]).split("_")
        (pre if kind == "prefix" else suf)[(t, a)] = str(r[5])
    deck = None
    spacing = None
    for r in wb["deck and spacing"].iter_rows(min_row=2, max_row=4, max_col=13, values_only=True):
        if r[0] is not None and str(r[0]).strip() == "78":
            deck = [float(x) for x in r[1:7]]
        if r[9] is not None and str(r[9]).startswith("Standard"):
            spacing = (float(r[10]), float(r[11]), float(r[12]))
    cached = [str(r[2]).strip() for r in wb["analysis"].iter_rows(min_row=7, max_row=23002, min_col=1, max_col=3, values_only=True)
              if isinstance(r[2], str) and r[2].strip() and r[2].strip() != "XXXXX"]
    sampler = None
    for r in wb["text dropdown"].iter_rows(min_row=2, values_only=True):
        if r[0] == "Four-Text Sampler":
            sampler = str(r[1])
    home = {str(r[0]): r[1] for r in wb["home"].iter_rows(min_row=1, max_row=21, max_col=2, values_only=True) if r[0]}
    assert len(uni) == 138 and len(pre) == 138 and len(suf) == 138, (len(uni), len(pre), len(suf))
    assert deck and spacing and sampler
    return dict(uni=uni, pre=pre, suf=suf, deck=deck, p_uni=spacing[0], spacing=spacing,
                cached=cached, sampler=sampler, home=home)


def synthetic_tables(rng):
    """Same shape as Greshko's (6 tables x 23 letters; unigram words; prefix and suffix pieces), random strings."""
    pfx_fam = ["qok", "ok", "qot", "ot", "olk", "sh"]
    sfx = ["aiin", "ain", "al", "ar", "edy", "eey", "ey", "y", "dy", "ol", "or", "am", "eedy", "eal", "ody", "oey",
           "aly", "air", "eol", "eody", "oly", "ory", "dal"]
    # synthetic prefix pieces each contain exactly one 'h', at their end, and no suffix contains 'h': uniquely splittable
    pieces = [b + "h" for b in ["c", "s", "k", "t", "p", "f", "d", "l", "r", "qoc", "oc", "okc", "otc", "lc", "dc", "tc",
                                "kc", "ckc", "ctc", "olc", "doc", "yc", "sc"]]
    uni, pre, suf = {}, {}, {}
    for ti, t in enumerate(TABLES):
        s = sfx[:]; rng.shuffle(s)
        p = pieces[:]; rng.shuffle(p)
        s2 = sfx[:]; rng.shuffle(s2)
        for li, a in enumerate(LETTERS23):
            uni[(t, a)] = pfx_fam[ti] + s[li]
            pre[(t, a)] = p[li]
            suf[(t, a)] = s2[li]
    return dict(uni=uni, pre=pre, suf=suf, deck=[0.358974358974359, 0.5384615384615384, 0.6794871794871795,
                                                0.8205128205128205, 0.9102564102564102, 1.0], p_uni=0.4722222222222222,
                spacing=(0.4722222222222222, 0.5277777777777778, 0.0), cached=None, sampler=None, home={})


def norm_letters(words):
    out = []
    for w in words:
        w = "".join(c for c in unicodedata.normalize("NFKD", str(w).lower()) if unicodedata.category(c) != "Mn")
        w = w.replace("j", "i").replace("k", "c").replace("w", "v")
        out.append("".join(c for c in w if c in LETTERS23))
    return "".join(out)


def draw_table(u, deck):
    for t, th in zip(TABLES, deck):
        if u < th:
            return t
    return TABLES[-1]


def encipher(letters, K, rng, n_tok):
    """workbook 'encrypter' procedure; returns tokens and the key record [(chunk, table1, table2)]."""
    toks, rec, i, cycled = [], [], 0, False
    L = len(letters)
    while len(toks) < n_tok:
        u = rng.random()
        k = 1 if u < K["p_uni"] else 2
        chunk = "".join(letters[(i + j) % L] for j in range(k))
        if i + k > L:
            cycled = True
        i = (i + k) % L if i + k >= L else i + k
        t1 = draw_table(rng.random(), K["deck"]); t2 = draw_table(rng.random(), K["deck"])
        if k == 1:
            toks.append(K["uni"][(t1, chunk)]); rec.append((chunk, t1, ""))
        else:
            toks.append(K["pre"][(t1, chunk[0])] + K["suf"][(t2, chunk[1])]); rec.append((chunk, t1, t2))
    return toks, rec, cycled


def decoders(K):
    inv_uni_t = {(t, w): a for (t, a), w in K["uni"].items()}
    inv_uni = {}
    for (t, a), w in K["uni"].items():
        inv_uni.setdefault(w, set()).add(a)
    inv_big = {}
    for (t1, a), p in K["pre"].items():
        for (t2, b), s in K["suf"].items():
            inv_big.setdefault(p + s, set()).add(a + b)
    return inv_uni_t, inv_uni, inv_big


def roundtrip(toks, rec, K, dec):
    inv_uni_t, inv_uni, inv_big = dec
    full_ok, len_unique, tab_unique = 0, 0, 0
    out = []
    for w, (chunk, t1, t2) in zip(toks, rec):
        if t2 == "":
            got = inv_uni_t.get((t1, w))
        else:
            got = None
            p_ = K["pre"]; s_ = K["suf"]
            for a in LETTERS23:
                pa = p_[(t1, a)]
                if not w.startswith(pa):
                    continue
                for b in LETTERS23:
                    if pa + s_[(t2, b)] == w:
                        got = a + b if got is None else (got if got == a + b else "?AMBIG?")
        out.append(got or "?")
        full_ok += got == chunk
        cand = inv_uni.get(w, set()) if t2 == "" else inv_big.get(w, set())
        len_unique += cand == {chunk}
        both = inv_uni.get(w, set()) | inv_big.get(w, set())
        tab_unique += both == {chunk}
    n = len(toks)
    return dict(exact="".join(out) == "".join(c for c, _, _ in rec), full_share=full_ok / n,
                len_unique=len_unique / n, tab_unique=tab_unique / n)


def repeat_rate(toks):
    t = toks[:10000]
    return sum(1 for a, b in zip(t, t[1:]) if a == b) / 10000.0


def mattr(toks, w=1000):
    if len(toks) < w:
        return float("nan")
    c = Counter(toks[:w]); tot = len(c); s = tot
    for i in range(w, len(toks)):
        o = toks[i - w]; c[o] -= 1
        if c[o] == 0:
            tot -= 1
        c[toks[i]] += 1
        if c[toks[i]] == 1:
            tot += 1
        s += tot
    return s / (len(toks) - w + 1) / w


def fidelity_row(name, toks, rec=None):
    r = dict(name=name, tokens=len(toks), types=len(set(toks)), repeat=repeat_rate(toks), mattr=mattr(toks),
             mean_len=float(np.mean([len(x) for x in toks])))
    r["uni_share"] = (sum(1 for c, _, _ in rec if len(c) == 1) / len(rec)) if rec else float("nan")
    return r


def pour(toks, skel_pages, n=None):
    it = iter(toks)
    out, used = [], 0
    for p in skel_pages:
        lines = []
        for ln in p["lines"]:
            k = len([w for w in ln if w])
            if k == 0:
                continue
            new = [w for _, w in zip(range(k), it)]
            used += len(new)
            if new:
                lines.append(new)
            if len(new) < k:
                break
        if lines:
            out.append(dict(id=str(p["id"]), section=str(p.get("section", "?")), lines=lines))
        if n is not None and used >= n:
            break
    return out, used


# ------------------------------------------------------------------ subsampling
def sub_f4(rows_by_page, ids, rng):
    arr = np.array([rows_by_page.get(i, np.zeros(3)) for i in ids])
    m = int(round(SUB_FRAC * len(ids)))
    return arr, m


def voy_subsample(td, B, rng):
    """S55.measure's bootstrap loop with 80 % of the held-out pages drawn WITHOUT replacement."""
    f1r = td.f1_rows(); gc, _, _ = td.gain_rows("core"); gp, _, _ = td.gain_rows("prefix")
    ids = [h["id"] for h in td.hp]
    F1 = np.array([f1r.get(i, np.zeros(6)) for i in ids]); GC = np.array([gc.get(i, np.zeros(3)) for i in ids])
    GP = np.array([gp.get(i, np.zeros(3)) for i in ids])
    n = len(td.hp); m = int(round(SUB_FRAC * n))
    out = np.full((B, len(FEATS)), np.nan)
    for b in range(B):
        ix = np.sort(rng.choice(n, m, replace=False))
        hps = [td.hp[i] for i in ix]
        c_, fr_, l_, s_, p_ = td.concat(hps)
        row = dict(f1_core_F1=S55.f1_stat(F1[ix]), f3_core_gain=S55.gain_stat(GC[ix]), f4_prefix_gain=S55.gain_stat(GP[ix]),
                   f2_core_core=td.core_core(c_, l_, N_SHUF_SUB, rng), f5a_core_cover=td.cover(hps),
                   f5b_core_heaps=S55.heaps(c_), f6_core_sec_nmi=td.sec_nmi(c_, s_, p_, N_SHUF_SUB, rng)[0],
                   f7_frame_core=td.frame_core(fr_, c_, N_SHUF_SUB, rng))
        out[b] = [row[k] for k in FEATS]
    return out


def cheap_sub(td, B, rng):
    """f1, f3, f4, f5a, f5b on 80 % held-out page subsamples (no shuffle nulls needed)."""
    f1r = td.f1_rows(); gc, _, _ = td.gain_rows("core"); gp, _, _ = td.gain_rows("prefix")
    ids = [h["id"] for h in td.hp]
    F1 = np.array([f1r.get(i, np.zeros(6)) for i in ids]); GC = np.array([gc.get(i, np.zeros(3)) for i in ids])
    GP = np.array([gp.get(i, np.zeros(3)) for i in ids])
    n = len(td.hp); m = int(round(SUB_FRAC * n))
    keys = ["f1_core_F1", "f3_core_gain", "f4_prefix_gain", "f5a_core_cover", "f5b_core_heaps"]
    out = np.full((B, len(keys)), np.nan)
    for b in range(B):
        ix = np.sort(rng.choice(n, m, replace=False))
        hps = [td.hp[i] for i in ix]
        out[b] = [S55.f1_stat(F1[ix]), S55.gain_stat(GC[ix]), S55.gain_stat(GP[ix]), td.cover(hps),
                  S55.heaps(np.concatenate([h["cores"] for h in hps]))]
    return keys, out, S55.gain_stat(GP)


# ------------------------------------------------------------------ main
def main():
    t0 = time.time()
    emit("S64: Naibbe cipher (Greshko 2025) in the s55 wrapper-free taxonomy, next to Voynich"
         + (" [SYNTHETIC FIXTURE, not a result]" if SYNTH else ""))
    emit(f"date: {S55.clock()}")
    emit(f"{SELF.name} sha256={sha256(SELF)}")
    emit(f"prereg stub {STUB.name}" + (f" sha256={sha256(STUB)}" if STUB.exists() else " (absent)") + f"  stub_match={stub_match()}")
    for m in S55.local_modules():
        emit(f"  local import {m} sha256={sha256(ROOT / m)}")
    import sklearn
    emit(f"numpy {np.__version__}; sklearn {sklearn.__version__}; python {sys.version.split()[0]}")
    emit(f"s55 constants: SEED {S55.SEED}, N_SHUF {S55.N_SHUF}, N_NULL_GAIN {S55.N_NULL_GAIN}; s64: N_TOK {N_TOK}, "
         f"subsample {SUB_FRAC:.0%} of held-out pages without replacement, B={B_SUB}, rng {SUB_SEED}, {N_SHUF_SUB} null draws per Voynich draw")
    emit()

    # ---------------- texts
    if SYNTH:
        texts, voy = S55.fixture_corpus()
        labels = OrderedDict((n, c) for n, (c, _) in texts.items() if c in S55.CLS)
        pages_of = {n: p for n, (c, p) in texts.items() if c in S55.CLS}
        voy_pages = voy
        K = synthetic_tables(random.Random(5))
        r0 = random.Random(9)
        plains = {}
        for nm in ["synA", "synB", "synC"]:
            # synthetic letter stream: first-order Markov over the 23 letters
            trans = {a: [r0.choice(LETTERS23) for _ in range(4)] for a in LETTERS23}
            s, buf = "a", []
            for _ in range(20000):
                s = r0.choice(trans[s]) if r0.random() < 0.8 else r0.choice(LETTERS23)
                buf.append(s)
            plains[nm] = "".join(buf)
        skel = {nm: pages_of["fx_code_0"] for nm in plains}
        n_tok = S55.TextData("tmp", pages_of["fx_code_0"]).tokens
        emit("synthetic: s55 fixture corpus (CODE 6 / GEN 6 / 2 VOY), synthetic Naibbe-shaped tables, Markov letter plaintexts")
    else:
        emit(f"Greshko workbook {XLSX.relative_to(ROOT).as_posix()} sha256={sha256(XLSX)} (read as data, values only)")
        K = load_greshko()
        emit(f"  tables: unigram {len(K['uni'])} entries, prefix {len(K['pre'])}, suffix {len(K['suf'])}; deck 78 thresholds "
             + ", ".join(f"{t} {x:.4f}" for t, x in zip(TABLES, K["deck"])) + f"; spacing Standard {K['spacing']}")
        emit(f"  S1 Table 2 spot check (alpha/beta1 a, i, t): {K['uni'][('alpha', 'a')]} {K['uni'][('beta1', 'a')]} "
             f"{K['uni'][('alpha', 'i')]} {K['uni'][('beta1', 'i')]} {K['uni'][('alpha', 't')]} {K['uni'][('beta1', 't')]}"
             "  (supplement: ol or / shedy qokedy / qokaiin okaiin)")
        lab_path = S49 / "labels.json"
        emit(f"labels: {lab_path.relative_to(ROOT).as_posix()} sha256={sha256(lab_path)}")
        lab49 = json.loads(lab_path.read_text(encoding="utf-8"))["labels"]
        labels = OrderedDict((n, c) for n, c in lab49.items() if c in S55.CLS)
        pages_of = {n: S55.load_sidecar(n) for n in labels}
        voy_pages = OrderedDict((f"Voynich_{T}", S55.voynich_pages(T)) for T in TRANSLITS)
        plains, skel = {}, {}
        for nm in PLAINS:
            p = S49 / f"lang_{nm}.meta.json"
            d = json.loads(p.read_text(encoding="utf-8"))
            plains[nm] = norm_letters([w for pg in d["pages"] for ln in pg["lines"] for w in ln])
            skel[nm] = S55.load_sidecar(f"code_k4_{nm}")
            emit(f"  plaintext {p.name} sha256={sha256(p)[:16]}: {len(plains[nm])} letters after normalisation; "
                 f"skeleton code_k4_{nm}.meta.json sha256={sha256(S49 / f'code_k4_{nm}.meta.json')[:16]}")
        n_tok = N_TOK
    emit()
    flush()

    # ---------------- Naibbe texts + round trip + fidelity
    emit("=== NAIBBE texts: encipherment, planted round-trip, fidelity ===")
    dec = decoders(K)
    amb_types = sum(1 for w in set(dec[1]) | set(dec[2]) if len(dec[1].get(w, set()) | dec[2].get(w, set())) > 1)
    emit(f"  table-only ambiguity: {amb_types} ciphertext types have more than one plaintext reading (unigram vs bigram or split)")
    naibbe = OrderedDict()
    fid = []
    rt_all = True
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    for pi, nm in enumerate(plains):
        for s in SEEDS:
            name = f"naibbe_{nm}_s{s}"
            rng = random.Random(64000 + 100 * pi + s)
            toks, rec, cyc = encipher(plains[nm], K, rng, n_tok)
            rt = roundtrip(toks, rec, K, dec)
            rt_all &= rt["exact"]
            pages, used = pour(toks, skel[nm])
            naibbe[name] = pages
            fr = fidelity_row(name, toks, rec); fid.append(fr)
            emit(f"  {name:30s} tokens {len(toks)} poured {used}; letters used {sum(len(c) for c, _, _ in rec)} of {len(plains[nm])}"
                 f"{' CYCLED' if cyc else ''}; round-trip with key record exact={rt['exact']} ({rt['full_share']:.4f}); "
                 f"unique reading given chunk lengths {rt['len_unique']:.4f}; tables only {rt['tab_unique']:.4f}")
            emit(f"      sample: {' '.join(toks[:12])}  <- {' '.join(c for c, _, _ in rec[:12])}")
            if not SYNTH:
                (RUN_DIR / f"{name}.meta.json").write_text(json.dumps(dict(
                    name=name, s64=dict(plaintext=nm, seed=64000 + 100 * pi + s, cycled=cyc,
                                        key_record="".join(f"{c}|{t1}|{t2};" for c, t1, t2 in rec)),
                    pages=pages), ensure_ascii=False), encoding="utf-8", newline="\n")
    extra = OrderedDict()
    if not SYNTH:
        cached = K["cached"]
        fid.insert(0, fidelity_row("greshko_cached (his output)", cached))
        for s in SEEDS:
            rng = random.Random(64900 + s)
            letters = norm_letters([K["sampler"]])
            toks, rec, cyc = encipher(letters, K, rng, len(cached))
            fid.insert(s, fidelity_row(f"s64 impl on his sampler s{s}", toks, rec))
        pages, used = pour(cached, skel["latin_cremma"], n=len(cached))
        extra["naibbe_greshko_cached"] = pages
        emit(f"  extra naibbe_greshko_cached: {len(cached)} tokens poured {used} (reported only, 21k tokens)")
    emit("  fidelity (repeat = workbook definition over first 10k tokens; MATTR = all 1000-token windows):")
    emit(f"    {'text':34s} {'tokens':>6s} {'types':>6s} {'repeat %':>8s} {'MATTR':>7s} {'mean len':>8s} {'uni share':>9s}")
    for r in fid:
        emit(f"    {r['name']:34s} {r['tokens']:6d} {r['types']:6d} {100 * r['repeat']:8.2f} {r['mattr']:7.4f} {r['mean_len']:8.3f} {r['uni_share']:9.4f}")
    if not SYNTH:
        h = K["home"]
        emit(f"    Greshko reported: repeat 0.18 +- 0.04 % (S1.6), MATTR 0.56 (S1.9); his workbook home sheet for the cached run: "
             f"repeat {float(h.get('% repeating pairs in the first 10,000 tokens', 'nan')) * 100:.2f} %, MATTR "
             f"{float(h.get('Average 1000-token MATTR, sampled 5 times', 'nan')):.4f}, types {h.get('Number of ciphertext word types')}")
    emit(f"  PLANTED round-trip (gating): {'PASS' if rt_all else 'FAIL'} on all {len(naibbe)} NAIBBE texts")
    emit()
    flush()

    # ---------------- features
    emit("=== features (s55 pipeline: S55.TextData(...).measure(); held-out pages) ===")
    emit(f"  {'text':30s} {'cls':6s} {'tok':>6s} {'teTok':>5s} {'tr/te':>7s} " + " ".join(f"{k[:14]:>14s}" for k in FEATS))
    F, TD = OrderedDict(), {}
    cls_of = OrderedDict()
    order = [(n, labels[n], pages_of[n]) for n in labels] + [(n, "NAIBBE", p) for n, p in naibbe.items()] \
        + [(n, "EXTRA", p) for n, p in extra.items()] + [(n, "VOY", p) for n, p in voy_pages.items()]
    for name, cls, pages in order:
        if time.time() - t0 > WALL:
            emit(f"STOPPED: wall limit {WALL} s at {S55.clock()}"); flush(); sys.exit(2)
        td = S55.TextData(name, pages)
        f, info, _ = td.measure(boot=False)
        F[name], TD[name], cls_of[name] = f, td, cls
        emit(f"  {name:30s} {cls:6s} {td.tokens:6d} {td.te_tokens:5d} {len(td.tr):3d}/{len(td.te):3d} "
             + " ".join(f"{f[k]:+14.4f}" for k in FEATS) + f"  [{time.time() - t0:.0f}s]")
        flush()
    if not SYNTH and S55_FEAT.exists():
        ref = {r["text"]: r for r in csv.DictReader(S55_FEAT.open(encoding="utf-8"))}
        diffs = [abs(F[n][k] - float(ref[n][k])) for n in F if n in ref for k in FEATS
                 if np.isfinite(F[n][k]) and ref[n][k] not in ("nan", "")]
        emit(f"  re-score check vs runs/s55/features.csv: {len([n for n in F if n in ref])} texts, max |diff| {max(diffs):.2e} "
             "(features.csv stores 6 significant digits)")
    with (RUN_DIR / ("fixture_features.csv" if SYNTH else "features.csv")).open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["text", "class"] + FEATS)
        for n, f in F.items():
            w.writerow([n, cls_of[n]] + [f"{f[k]:.6g}" for k in FEATS])
    emit()

    core12 = list(labels)
    nb = list(naibbe)
    vn = list(voy_pages)

    # ---------------- table of class means
    emit("=== per-feature: s55 classes, NAIBBE (9 texts), Voynich ===")
    emit(f"  {'feature':16s} {'CODE mean (sd)':>19s} {'GEN mean (sd)':>19s} {'NAIBBE mean (sd)':>19s} {'NAIBBE min..max':>17s}   "
         + " ".join(f"{v:>13s}" for v in vn) + ("   greshko_cached" if extra else ""))
    for k in FEATS:
        a = [F[n][k] for n in core12 if labels[n] == "CODE"]; g = [F[n][k] for n in core12 if labels[n] == "GEN"]
        z = [F[n][k] for n in nb]
        emit(f"  {k:16s} {np.nanmean(a):+9.4f} ({np.nanstd(a):.4f}) {np.nanmean(g):+9.4f} ({np.nanstd(g):.4f}) "
             f"{np.nanmean(z):+9.4f} ({np.nanstd(z):.4f}) {np.nanmin(z):+.4f}..{np.nanmax(z):+.4f}   "
             + " ".join(f"{F[v][k]:+13.4f}" for v in vn) + (f"   {F['naibbe_greshko_cached'][k]:+.4f}" if extra else ""))
    emit()
    flush()

    # ---------------- P1
    emit(f"=== P1: f4_prefix_gain, NAIBBE below Voynich (80 % held-out page subsamples, without replacement, B={B_SUB}) ===")
    rng = np.random.default_rng(SUB_SEED)
    SUBK = {}
    per_hi = {}
    cls_draws = np.zeros(B_SUB)
    with (RUN_DIR / ("fixture_naibbe_subsample.csv" if SYNTH else "naibbe_subsample.csv")).open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        for n in nb:
            keys, arr, pt = cheap_sub(TD[n], B_SUB, rng)
            assert abs(pt - F[n]["f4_prefix_gain"]) < 1e-12, (n, pt, F[n]["f4_prefix_gain"])
            SUBK[n] = (keys, arr)
            if not fh.tell():
                w.writerow(["text", "draw"] + keys)
            for bi, row in enumerate(arr):
                w.writerow([n, bi] + [f"{x:.6g}" for x in row])
            col = arr[:, keys.index("f4_prefix_gain")]
            cls_draws += col / len(nb)
            lo, hi = np.percentile(col, [2.5, 97.5]); per_hi[n] = hi
            emit(f"  {n:30s} f4 {F[n]['f4_prefix_gain']:+.4f} [{lo:+.4f}, {hi:+.4f}]")
    clo, chi = np.percentile(cls_draws, [2.5, 97.5])
    cmean = float(np.mean([F[n]["f4_prefix_gain"] for n in nb]))
    emit(f"  NAIBBE class mean f4 {cmean:+.4f} [{clo:+.4f}, {chi:+.4f}]; largest per-text 97.5th pct {max(per_hi.values()):+.4f}")
    emit("  other features, NAIBBE class-mean 80 %-subsample intervals (reported): "
         + "; ".join(f"{k} [{np.percentile(np.mean([SUBK[n][1][:, i] for n in nb], 0), 2.5):+.4f}, "
                     f"{np.percentile(np.mean([SUBK[n][1][:, i] for n in nb], 0), 97.5):+.4f}]"
                     for i, k in enumerate(SUBK[nb[0]][0])))
    P1 = {}
    for v in vn:
        fv = F[v]["f4_prefix_gain"]
        P1[v] = bool(chi < fv and max(per_hi.values()) < fv)
        emit(f"  {v}: f4 {fv:+.4f}; class 97.5th {chi:+.4f} < {fv:+.4f}: {chi < fv}; all per-text 97.5th < {fv:+.4f}: "
             f"{max(per_hi.values()) < fv} -> P1 {'HOLDS' if P1[v] else 'FAILS'}")
    emit()
    flush()

    # ---------------- P2
    emit("=== P2: Voynich outside the NAIBBE region (s55 z-space; centroids CODE, GEN, NAIBBE; LOO envelope) ===")
    X12 = np.array([[F[n][k] for k in FEATS] for n in core12], dtype=float)
    y = [labels[n] for n in core12]
    Zall = S55.prep(X12, np.array([[F[n][k] for k in FEATS] for n in nb + vn + list(extra)], dtype=float))
    Z12, Zo = Zall
    ZN = Zo[:len(nb)]; ZV = Zo[len(nb):len(nb) + len(vn)]; ZE = Zo[len(nb) + len(vn):]
    cents = {c: Z12[[i for i, t in enumerate(y) if t == c]].mean(0) for c in S55.CLS}
    cents["NAIBBE"] = ZN.mean(0)
    loo_d = [float(np.linalg.norm(ZN[i] - np.delete(ZN, i, 0).mean(0))) for i in range(len(nb))]
    R = max(loo_d)
    emit(f"  NAIBBE LOO distances to the centroid of the other 8: " + ", ".join(f"{d:.3f}" for d in loo_d) + f"  -> R = {R:.3f}")
    emit("  centroid distances (z): " + "; ".join(f"{a}-{b} {np.linalg.norm(cents[a] - cents[b]):.3f}"
                                               for a, b in (("CODE", "GEN"), ("CODE", "NAIBBE"), ("GEN", "NAIBBE"))))
    emit("  NAIBBE z centroid: " + ", ".join(f"{k} {z:+.2f}" for k, z in zip(FEATS, cents["NAIBBE"])))
    m_lr = S55.lr_fit(Z12, y)
    for n, z, lp in zip(nb, ZN, S55.lr_predict(m_lr, ZN)):
        d = {c: float(np.linalg.norm(z - m)) for c, m in cents.items()}
        emit(f"  {n:30s} nearest of 3 {min(d, key=d.get):6s} (d CODE {d['CODE']:.3f}, GEN {d['GEN']:.3f}, NAIBBE {d['NAIBBE']:.3f}); "
             f"s55 LR (2-class) {lp[0]} P(CODE) {lp[2]:.3f}")
    for n, z in zip(extra, ZE):
        d = {c: float(np.linalg.norm(z - m)) for c, m in cents.items()}
        emit(f"  {n:30s} [extra] nearest of 3 {min(d, key=d.get)} (d CODE {d['CODE']:.3f}, GEN {d['GEN']:.3f}, NAIBBE {d['NAIBBE']:.3f})")
    P2 = {}
    vrng = np.random.default_rng(SUB_SEED + 1)
    for v, z in zip(vn, ZV):
        d = {c: float(np.linalg.norm(z - m)) for c, m in cents.items()}
        near = min(d, key=d.get)
        P2[v] = bool(near != "NAIBBE" and d["NAIBBE"] > R)
        emit(f"  {v}: nearest of 3 {near} (d CODE {d['CODE']:.3f}, GEN {d['GEN']:.3f}, NAIBBE {d['NAIBBE']:.3f}); "
             f"d NAIBBE {d['NAIBBE']:.3f} > R {R:.3f}: {d['NAIBBE'] > R} -> P2 {'HOLDS' if P2[v] else 'FAILS'}")
        emit(f"      z-vector: " + ", ".join(f"{k} {q:+.2f}" for k, q in zip(FEATS, z)))
        if time.time() - t0 < WALL - 300:
            Xb = voy_subsample(TD[v], B_SUB, vrng)
            Zb = S55.prep(X12, Xb)[1]
            ok = 0
            for zb in Zb:
                db = {c: float(np.linalg.norm(zb - m)) for c, m in cents.items()}
                ok += (min(db, key=db.get) != "NAIBBE") and db["NAIBBE"] > R
            f4c = Xb[:, FEATS.index("f4_prefix_gain")]
            emit(f"      Voynich 80 % subsamples (B={B_SUB}, reported): share meeting P2 {ok / B_SUB:.3f}; f4 "
                 f"[{np.nanpercentile(f4c, 2.5):+.4f}, {np.nanpercentile(f4c, 97.5):+.4f}]; "
                 + "; ".join(f"{k} [{np.nanpercentile(Xb[:, i], 2.5):+.4f}, {np.nanpercentile(Xb[:, i], 97.5):+.4f}]"
                             for i, k in enumerate(FEATS) if k != "f4_prefix_gain"))
            with (RUN_DIR / ("fixture_voynich_subsample.csv" if SYNTH else "voynich_subsample.csv")).open(
                    "a" if v != vn[0] else "w", encoding="utf-8", newline="") as fh:
                w = csv.writer(fh, lineterminator="\n")
                if v == vn[0]:
                    w.writerow(["text", "draw"] + FEATS)
                for bi, row in enumerate(Xb):
                    w.writerow([v, bi] + [f"{x:.6g}" for x in row])
        flush()
    emit()

    # ---------------- verdict
    emit("=== READING (frozen rules) ===")
    for v in vn:
        emit(f"  {v}: P1 {'HOLDS' if P1[v] else 'FAILS'}; P2 {'HOLDS' if P2[v] else 'FAILS'}")
    if not rt_all:
        verdict = "INVALID (implementation does not round-trip)"
    elif all(P1[v] and P2[v] for v in vn):
        verdict = "random-table-per-token cipher excluded"
    else:
        verdict = "not excluded"
    emit(f"VERDICT|s64|{verdict}|roundtrip={'PASS' if rt_all else 'FAIL'}|"
         + "|".join(f"{v}=P1:{'Y' if P1[v] else 'N'},P2:{'Y' if P2[v] else 'N'}" for v in vn))
    emit()
    emit(f"elapsed {time.time() - t0:.0f}s; run ended {S55.clock()}")
    emit(f"outputs: {OUT}; {RUN_DIR}")
    emit(f"stub_match at end (all local modules loaded by now): {stub_match()}")
    flush()


if __name__ == "__main__":
    main()
